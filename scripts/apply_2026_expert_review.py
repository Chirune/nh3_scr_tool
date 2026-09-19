"""Apply the five grouped expert decisions for the 2026 Cu-CHA case."""
import argparse
import copy
import json
from pathlib import Path

from openpyxl import load_workbook

from scrtool.cli import run_export
from scrtool.core import FIELDS, quantity, uid, write_csv, write_json
from scrtool.features import build_features


def read_decisions(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["需要你判断"]
    rows = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if values[0]:
            rows[str(values[0])] = {"choice": values[4], "note": values[5]}
    workbook.close()
    return rows


def approve(record, decision, note):
    record["issues"] = [issue for issue in record.get("issues", []) if issue != "cross_source_conflict"]
    record["review_status"] = "approved"
    record["reviewer"] = "User domain expert"
    record["review_level"] = "expert_reviewed"
    record["review_note"] = note
    record["expert_decision_id"] = decision


def reject(record, decision, note):
    record["review_status"] = "rejected"
    record["reviewer"] = "User domain expert"
    record["review_level"] = "expert_reviewed"
    record["review_note"] = note
    record["expert_decision_id"] = decision


def clone_quantity(seed, decision, catalyst, prop, raw, raw_unit, note):
    record = copy.deepcopy(seed)
    value, issue = quantity(raw, raw_unit, prop)
    record.update(
        record_id=uid("expert_resolution", decision, catalyst, prop, raw),
        catalyst=catalyst,
        property=prop,
        category=FIELDS[prop][0],
        raw_value=str(raw),
        raw_unit=raw_unit,
        value=value,
        unit=FIELDS[prop][1],
        issues=[issue] if issue else [],
        method="expert_resolution",
        estimated=False,
    )
    approve(record, decision, note)
    return record


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--review", type=Path, required=True)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    decisions = read_decisions(args.review)
    expected = {"D001", "D002", "D003", "D004", "D005"}
    if set(decisions) != expected or any(decisions[key]["choice"] in (None, "") for key in expected):
        raise ValueError("All five review decisions must be filled")

    chosen = {"Cu-0.5": float(decisions["D001"]["choice"]), "Cu-1.4": float(decisions["D002"]["choice"])}
    records = json.loads(args.input.read_text(encoding="utf-8"))
    additions = []

    for record in records:
        if record.get("review_status") != "pending":
            continue
        catalyst, prop = record.get("catalyst"), record.get("property")
        if catalyst in chosen and prop == "cu_content":
            if record.get("value") == chosen[catalyst]:
                approve(record, "D001" if catalyst == "Cu-0.5" else "D002",
                        f"Expert selected {chosen[catalyst]:.2f} wt% as the modeling value; alternate publication rounding retained as rejected provenance.")
            else:
                reject(record, "D001" if catalyst == "Cu-0.5" else "D002",
                       f"Expert selected {chosen[catalyst]:.2f} wt% for modeling.")

    temperature = next(r for r in records if r.get("review_status") == "pending" and r.get("property") == "temperature" and not r.get("catalyst"))
    reject(temperature, "D003", "Replaced by two sample-linked condition records.")
    for catalyst in ("CMI-DG *", "CMI-FA **"):
        additions.append(clone_quantity(temperature, "D003", catalyst, "temperature", "200", "°C",
                                        "Expert confirmed 200 °C is shared by the DG and FA series in Supplementary Fig. 1."))

    for record in records:
        if record.get("review_status") == "pending" and record.get("raw_value") == "--":
            reject(record, "D004", "Expert confirmed '--' means not reported; it is not a numeric zero.")

    composite = [r for r in records if r.get("review_status") == "pending" and r.get("catalyst") == "CMI-FA **" and "(" in str(r.get("raw_value"))]
    seeds = {r["property"]: r for r in composite}
    for record in composite:
        reject(record, "D005", "Replaced by total-Cu and isolated-Cu records according to Supplementary Table 1 footnote a.")
    split_specs = [
        ("cu_content", "cu_content", "2.42", "wt%"),
        ("cu_content", "isolated_cu_content", "1.39", "wt%"),
        ("cu_al_ratio", "cu_al_ratio", "0.30", "1"),
        ("cu_al_ratio", "isolated_cu_al_ratio", "0.18", "1"),
        ("cu_per_cage", "cu_per_cage", "0.29", "1"),
        ("cu_per_cage", "isolated_cu_per_cage", "0.17", "1"),
    ]
    for seed_prop, prop, raw, unit in split_specs:
        additions.append(clone_quantity(seeds[seed_prop], "D005", "CMI-FA **", prop, raw, unit,
                                        "Split using Supplementary Table 1 footnote a; isolated fields are EPR-based."))

    unresolved = [r["record_id"] for r in records if r.get("review_status") == "pending"]
    if unresolved:
        raise ValueError("Pending records were not resolved: " + ", ".join(unresolved))
    final = records + additions
    args.output.mkdir(parents=True, exist_ok=True)
    write_json(args.output / "combined_expert_reviewed.json", final)
    write_json(args.output / "expert_decisions.json", decisions)
    write_csv(args.output / "all_observations.csv", final)
    write_csv(args.output / "rejected.csv", [r for r in final if r.get("review_status") == "rejected"])
    run_export(argparse.Namespace(input=str(args.output / "combined_expert_reviewed.json"), output=str(args.output / "training")))
    build_features(final, args.output / "training")
    summary = {
        "input_records": len(records),
        "added_resolved_records": len(additions),
        "final_records": len(final),
        "approved": sum(r.get("review_status") == "approved" for r in final),
        "rejected": sum(r.get("review_status") == "rejected" for r in final),
        "pending": sum(r.get("review_status") == "pending" for r in final),
        "expert_decisions": decisions,
    }
    write_json(args.output / "expert_review_report.json", summary)
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()

"""Finalize the reviewed 2020 gold set and measure automatic field coverage.

This evaluates whether each expert-approved gold item was recovered.  It does
not call unlisted automatic candidates false positives because this compact
gold set is a coverage checklist rather than an exhaustive annotation of every
valid observation in the paper.
"""
import argparse
import csv
import json
import re
from pathlib import Path


SUBSCRIPTS = str.maketrans("₀₁₂₃₄₅₆₇₈₉", "0123456789")


def text(value):
    return str(value or "").translate(SUBSCRIPTS).replace("α", "alpha").strip().lower()


def number(value):
    cleaned = re.sub(r"(?<=\d),(?=\d{3}(?:\D|$))", "", str(value or "").replace("−", "-"))
    match = re.search(r"[-+]?\d+(?:\.\d+)?(?:e[-+]?\d+)?", cleaned, re.I)
    return float(match.group()) if match else None


def unit(value):
    return text(value).replace("vol%", "%")


def read_gold(path):
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook["审核表"] if "审核表" in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(v or "") for v in next(rows)]
        result = [dict(zip(headers, values)) for values in rows if any(v is not None for v in values)]
        workbook.close()
        return result
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def read_csv(path):
    if not path.exists():
        return []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def scalar_matches(gold, candidates):
    aliases = {"mo_loading": {"mo_loading", "loading"}}
    fields = aliases.get(text(gold["field"]), {text(gold["field"])})
    target_number = number(gold.get("value"))
    matches = []
    for candidate in candidates:
        if text(candidate.get("property")) not in fields:
            continue
        candidate_number = number(candidate.get("raw_value") or candidate.get("value"))
        if target_number is not None and candidate_number != target_number:
            continue
        oxidation_unit = text(gold.get("field")) == "mo_oxidation_state" and text(gold.get("unit")) in {"+", "1"}
        if gold.get("unit") and not oxidation_unit and unit(candidate.get("unit") or candidate.get("raw_unit")) != unit(gold.get("unit")):
            continue
        matches.append(candidate)
    strict = [c for c in matches if text(c.get("catalyst")) == text(gold.get("catalyst"))]
    return strict, matches


def curve_match(field, curve_root):
    files = {
        "no_conversion_curve": curve_root / "fig11" / "points.csv",
        "n2_selectivity_curve": curve_root / "fig12" / "points.csv",
        "durability_curve": curve_root / "fig13" / "points.csv",
    }
    path = files.get(field)
    rows = read_csv(path) if path else []
    if field == "no_conversion_curve":
        rows = [r for r in rows if r.get("y_property") == "no_conversion"]
    elif field == "n2_selectivity_curve":
        rows = [r for r in rows if r.get("y_property") == "n2_selectivity"]
    elif field == "durability_curve":
        rows = [r for r in rows if r.get("x_property") == "time_on_stream" and r.get("y_property") == "no_conversion"]
    return path, rows


def write_csv(path, rows, headers):
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--gold", type=Path, required=True)
    parser.add_argument("--main", type=Path, required=True)
    parser.add_argument("--supp", type=Path, required=True)
    parser.add_argument("--curves", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)

    gold = read_gold(args.gold)
    invalid = [r.get("row_id") for r in gold if text(r.get("expert_decision")) not in {"approve", "reject", "uncertain"}]
    if invalid:
        raise ValueError("Unreviewed or invalid gold decisions: " + ", ".join(map(str, invalid)))
    approved = [r for r in gold if text(r.get("expert_decision")) == "approve"]
    candidates = read_csv(args.main) + read_csv(args.supp)
    coverage = []
    for row in approved:
        field = text(row.get("field"))
        if field.endswith("_curve"):
            path, points = curve_match(field, args.curves)
            status = "covered" if points else "missing"
            detail = f"{len(points)} vector-marker points in {path}" if points else "curve output not found"
        else:
            strict, loose = scalar_matches(row, candidates)
            if strict:
                status = "covered"
                detail = f"{len(strict)} strict candidate match(es)"
            elif loose:
                status = "partial"
                detail = "value found, but catalyst/sample identity does not match or is missing"
            else:
                status = "missing"
                detail = "no matching automatic candidate"
        coverage.append({**row, "coverage_status": status, "coverage_detail": detail})

    counts = {key: sum(r["coverage_status"] == key for r in coverage) for key in ("covered", "partial", "missing")}
    report = {
        "paper_id": "10.1038/s41467-020-15261-5",
        "approved_gold_items": len(approved),
        **counts,
        "strict_gold_coverage": counts["covered"] / len(approved) if approved else None,
        "precision": None,
        "precision_note": "Not computed: this gold set is not an exhaustive list of all valid observations.",
        "curve_point_accuracy": None,
        "curve_point_accuracy_note": "G022-G024 approve curve sources/categories, not each digitized numeric point.",
    }
    headers = list(gold[0]) if gold else []
    write_csv(args.output / "gold_2020_approved.csv", approved, headers)
    write_csv(args.output / "gold_2020_coverage.csv", coverage, headers + ["coverage_status", "coverage_detail"])
    (args.output / "gold_2020_approved.json").write_text(json.dumps(approved, ensure_ascii=False, indent=2), encoding="utf-8")
    (args.output / "evaluation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = [
        "# 2020 论文提取覆盖评估",
        "",
        f"- 人工批准金标准：{len(approved)} 项",
        f"- 严格自动覆盖：{counts['covered']} 项（{report['strict_gold_coverage']:.1%}）",
        f"- 部分覆盖：{counts['partial']} 项",
        f"- 未覆盖：{counts['missing']} 项",
        "- 精确率：未计算；当前金标准不是论文所有有效观测的穷尽标注。",
        "- 曲线点准确率：未计算；G022-G024 只确认曲线类别和来源，不代表逐点数值审核。",
        "",
        "## 逐项结果",
        "",
        "| 编号 | 字段 | 状态 | 说明 |",
        "|---|---|---|---|",
    ]
    lines.extend(f"| {r['row_id']} | {r['field']} | {r['coverage_status']} | {r['coverage_detail']} |" for r in coverage)
    (args.output / "evaluation.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()

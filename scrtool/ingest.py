import csv
import hashlib
import json
import re
import subprocess
import os
import sys
from html.parser import HTMLParser
from pathlib import Path
from .core import uid, read_json


class _ArticleHTMLParser(HTMLParser):
    """Extract visible article text and simple tables from saved publisher HTML."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.meta = {}
        self.blocks = []
        self._hidden = 0
        self._capture = None
        self._parts = []
        self._table = []
        self._row = []
        self._cell = []
        self._in_cell = False
        self._cell_rowspan = 1
        self._cell_colspan = 1
        self._spans = {}
        self._column = 0

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag in {'script', 'style', 'noscript', 'svg'}:
            self._hidden += 1
            return
        if self._hidden:
            return
        if tag == 'meta':
            key = (attrs.get('name') or attrs.get('property') or '').lower()
            if key and attrs.get('content'):
                self.meta[key] = attrs['content'].strip()
        if tag in {'title', 'h1', 'h2', 'h3', 'p', 'li', 'figcaption'}:
            self._capture = tag
            self._parts = []
        elif tag == 'table':
            self._table = []
            self._spans = {}
        elif tag == 'tr':
            self._row = []
            self._column = 0
        elif tag in {'th', 'td'}:
            while self._column in self._spans:
                remaining, value = self._spans[self._column]
                self._row.append(value)
                if remaining <= 1:
                    del self._spans[self._column]
                else:
                    self._spans[self._column] = (remaining - 1, value)
                self._column += 1
            self._cell = []
            self._in_cell = True
            try:
                self._cell_rowspan = max(1, int(attrs.get('rowspan', 1)))
                self._cell_colspan = max(1, int(attrs.get('colspan', 1)))
            except ValueError:
                self._cell_rowspan = self._cell_colspan = 1

    def handle_endtag(self, tag):
        if tag in {'script', 'style', 'noscript', 'svg'}:
            self._hidden = max(0, self._hidden - 1)
            return
        if self._hidden:
            return
        if tag in {'title', 'h1', 'h2', 'h3', 'p', 'li', 'figcaption'} and self._capture == tag:
            text = re.sub(r'\s+', ' ', ''.join(self._parts)).strip()
            if text:
                self.blocks.append((tag, text))
            self._capture = None
            self._parts = []
        elif tag in {'th', 'td'}:
            value = re.sub(r'\s+', ' ', ''.join(self._cell)).strip()
            for _ in range(self._cell_colspan):
                self._row.append(value)
                if self._cell_rowspan > 1:
                    self._spans[self._column] = (self._cell_rowspan - 1, value)
                self._column += 1
            self._cell = []
            self._in_cell = False
        elif tag == 'tr' and self._row:
            while self._column in self._spans:
                remaining, value = self._spans[self._column]
                self._row.append(value)
                if remaining <= 1:
                    del self._spans[self._column]
                else:
                    self._spans[self._column] = (remaining - 1, value)
                self._column += 1
            self._table.append(self._row)
            self._row = []
        elif tag == 'table' and self._table:
            self.blocks.append(('html_table', self._table))
            self._table = []

    def handle_data(self, data):
        if self._hidden:
            return
        if self._capture is not None:
            self._parts.append(data)
        if self._in_cell:
            self._cell.append(data)


def blocks(path, paper_id=None):
    path = Path(path).resolve()
    source_id = hashlib.sha256(path.read_bytes()).hexdigest()
    paper_id = paper_id or source_id[:20]

    def block(text, locator, kind='text', **extra):
        return dict(source_id=source_id, paper_id=paper_id, source_file=str(path),
                    text=text, locator=locator, kind=kind,
                    block_id=uid(source_id, locator, text), **extra)

    suffix = path.suffix.lower()
    if suffix == '.pdf':
        from pypdf import PdfReader
        for i, page in enumerate(PdfReader(path).pages, 1):
            text = page.extract_text() or ''
            yield block(text, f'page:{i}', 'pdf_text', needs_ocr=len(text.strip()) < 50)
        # Extract ruled tables separately. Merged-cell propagation uses actual
        # geometry; ordinary blank cells are never blindly forward-filled.
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            for pn, page in enumerate(pdf.pages, 1):
                for tn, table in enumerate(page.find_tables(), 1):
                    data = table.extract()
                    if len(data) < 2 or len(data[0]) < 2:
                        continue
                    headers = [re.sub(r'\s+', ' ', h or '').strip() for h in data[0]]
                    if not all(headers) or len(set(headers)) != len(headers):
                        continue
                    for ri, values in enumerate(data[1:], 1):
                        merged = []
                        for ci, value in enumerate(values):
                            if value is not None:
                                continue
                            header_cell = table.rows[0].cells[ci]
                            current_row = table.rows[ri]
                            if header_cell is None:
                                continue
                            x = (header_cell[0] + header_cell[2]) / 2
                            y = (current_row.bbox[1] + current_row.bbox[3]) / 2
                            spanning = next((cell for cell in table.cells if cell[0] < x < cell[2] and cell[1] < y < cell[3]), None)
                            if spanning:
                                values[ci] = page.crop(spanning).extract_text() or None
                                merged.append(headers[ci])
                        row = dict(zip(headers, values))
                        yield block(json.dumps(row, ensure_ascii=False), f'page:{pn};table:{tn};row:{ri}', 'pdf_table', row=row, bbox=list(table.rows[ri].bbox), merged_columns=merged)
    elif suffix in ['.csv', '.tsv']:
        with path.open(encoding='utf-8-sig', newline='') as f:
            reader = csv.DictReader(f, delimiter='\t' if suffix == '.tsv' else ',')
            if not reader.fieldnames or len(set(reader.fieldnames)) != len(reader.fieldnames):
                raise ValueError('Missing or duplicate CSV headers')
            for i, row in enumerate(reader, 2):
                if None in row or None in row.values():
                    raise ValueError(f'Malformed CSV row {i}')
                kind = 'digitized_curve' if str(row.get('estimated', '')).lower() == 'true' else 'raw_table'
                yield block(json.dumps(row, ensure_ascii=False), f'row:{i}', kind, row=row)
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in wb:
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, ())
                names = [str(h) if h is not None else f'unnamed_{i}' for i, h in enumerate(headers)]
                if len(set(names)) != len(names):
                    raise ValueError('Duplicate Excel column names: ' + sheet.title)
                for i, values in enumerate(rows, 2):
                    row = dict(zip(names, values))
                    yield block(json.dumps(row, ensure_ascii=False, default=str), f'sheet:{sheet.title};row:{i}', 'raw_table', row=row)
        finally:
            wb.close()
    elif suffix == '.json':
        content = read_json(path)
        if not isinstance(content, list):
            raise ValueError('JSON input must be a MinerU content_list (list of blocks)')
        for i, item in enumerate(content):
            if not isinstance(item, dict):
                raise ValueError('Invalid MinerU block')
            kind = item.get('type', 'text')
            text = '\n'.join(str(item.get(k, '')) for k in ['text', 'table_body', 'table_caption', 'image_caption', 'table_footnote'])
            table_html = str(item.get('table_body', ''))
            if kind == 'table' and '<table' in table_html.lower():
                parser = _ArticleHTMLParser()
                parser.feed(table_html)
                tables = [value for block_kind, value in parser.blocks if block_kind == 'html_table']
                for table_number, data in enumerate(tables, 1):
                    if len(data) < 2:
                        continue
                    headers = [re.sub(r'\s+', ' ', h or '').strip() for h in data[0]]
                    if not all(headers) or len(set(headers)) != len(headers):
                        yield block(text, f"page:{int(item.get('page_idx', 0)) + 1};item:{i};table:{table_number}", 'mineru_table_unresolved', image_path=item.get('img_path'))
                        continue
                    for row_number, values in enumerate(data[1:], 1):
                        if len(values) != len(headers):
                            yield block(text, f"page:{int(item.get('page_idx', 0)) + 1};item:{i};table:{table_number};row:{row_number}", 'mineru_table_unresolved', image_path=item.get('img_path'))
                            continue
                        row = dict(zip(headers, values))
                        yield block(json.dumps(row, ensure_ascii=False), f"page:{int(item.get('page_idx', 0)) + 1};item:{i};table:{table_number};row:{row_number}", 'mineru_table', row=row, image_path=item.get('img_path'))
                continue
            yield block(text, f"page:{int(item.get('page_idx', 0)) + 1};item:{i}", kind, image_path=item.get('img_path'))
    elif suffix in ['.html', '.htm']:
        parser = _ArticleHTMLParser()
        parser.feed(path.read_text(encoding='utf-8-sig', errors='replace'))
        meta = parser.meta
        for i, (kind, content) in enumerate(parser.blocks, 1):
            if kind == 'html_table':
                if len(content) < 2:
                    continue
                headers = [re.sub(r'\s+', ' ', str(h)).strip() for h in content[0]]
                if not all(headers) or len(set(headers)) != len(headers):
                    # Preserve malformed/merged tables for review rather than guessing cells.
                    yield block(json.dumps(content, ensure_ascii=False), f'html:table:{i}', 'html_table_unresolved', document_meta=meta)
                    continue
                for ri, values in enumerate(content[1:], 1):
                    if len(values) != len(headers):
                        yield block(json.dumps(values, ensure_ascii=False), f'html:table:{i};row:{ri}', 'html_table_unresolved', document_meta=meta)
                        continue
                    row = dict(zip(headers, values))
                    yield block(json.dumps(row, ensure_ascii=False), f'html:table:{i};row:{ri}', 'html_table', row=row, document_meta=meta)
            else:
                yield block(content, f'html:{kind}:{i}', 'html_text', document_meta=meta)
    elif suffix in ['.md', '.txt']:
        lines = path.read_text(encoding='utf-8-sig').splitlines()
        i = 0
        while i < len(lines):
            # Markdown tables become individual rows with explicit column context.
            if i + 1 < len(lines) and '|' in lines[i] and re.fullmatch(r'[\s|:\-]+', lines[i + 1]):
                headers = [s.strip() for s in lines[i].strip().strip('|').split('|')]
                i += 2
                while i < len(lines) and '|' in lines[i]:
                    values = [s.strip() for s in lines[i].strip().strip('|').split('|')]
                    if len(values) != len(headers) or len(set(headers)) != len(headers):
                        raise ValueError(f'Malformed Markdown table at line {i + 1}')
                    row = dict(zip(headers, values))
                    yield block(json.dumps(row, ensure_ascii=False), f'line:{i + 1}', 'table', row=row)
                    i += 1
                continue
            start = i
            chunk = []
            while i < len(lines) and (i == start or (lines[i].strip() and not ('|' in lines[i] and i + 1 < len(lines) and re.fullmatch(r'[\s|:\-]+', lines[i + 1])))):
                chunk.append(lines[i])
                i += 1
                if sum(map(len, chunk)) > 10000:
                    break
            text = '\n'.join(chunk)
            if text.strip():
                yield block(text, f'lines:{start + 1}-{i}', 'text')
    else:
        raise ValueError('Unsupported input: ' + suffix)


def convert(path, output, backend, executable=None, formula=True, tables=True, device='cpu'):
    """Use installed CLI entrypoints; avoid importing incompatible ML stacks."""
    Path(output).mkdir(parents=True, exist_ok=True)
    if backend == 'mineru':
        local_exe = Path(sys.executable).parent / ('mineru.exe' if os.name == 'nt' else 'mineru')
        cmd = [executable or (str(local_exe) if local_exe.exists() else 'mineru'), '-p', str(Path(path).resolve()), '-o', str(Path(output).resolve()), '-b', 'pipeline', '-f', str(formula).lower(), '-t', str(tables).lower()]
    else:
        cmd = [executable or 'marker_single', str(Path(path).resolve()), '--output_dir', str(Path(output).resolve())]
    env = os.environ.copy()
    project = Path(__file__).resolve().parents[1]
    env.setdefault('MINERU_DEVICE_MODE', device)
    env.setdefault('MINERU_MODEL_SOURCE', 'modelscope')
    env.setdefault('MINERU_TOOLS_CONFIG_JSON', str(project / 'mineru.json'))
    env.setdefault('MODELSCOPE_CACHE', str(project / '.models/modelscope'))
    env.setdefault('HF_HOME', str(project / '.models/huggingface'))
    subprocess.run(cmd, check=True, timeout=7200, env=env)
